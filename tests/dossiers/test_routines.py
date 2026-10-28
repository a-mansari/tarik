"""Routines: their schedule, made only by those who may (on every dossier Tarik sees) and on a yes, each period run
once, their report sent to a live collaborator, and a routine's run never confirming nor writing unattended."""

import asyncio
import io
from datetime import datetime, timedelta

import openpyxl
import pytest

from engine import calls, graphs
from engine import routine as flow
from engine import routines as clock
from engine.tools.core.cabinets import cabinet_db
from engine.tools.dossiers import routine_report as report_files
from engine.tools.dossiers import routines
from engine.tools.dossiers.routines import ZONE, next_run
from engine.tools.emails.sending import inline_html, plain, render


def at(text):
    return datetime.fromisoformat(text).replace(tzinfo=ZONE)


@pytest.mark.functionality
def test_the_next_time_a_schedule_falls_on():
    assert next_run("day", [], 8, 0, at("2026-10-27 07:59")) == at("2026-10-27 08:00")
    assert next_run("day", [], 8, 0, at("2026-10-27 08:00")) == at("2026-10-28 08:00")  # its own time: the next one
    assert next_run("week", [0], 9, 0, at("2026-10-27 10:00")) == at("2026-11-02 09:00")  # a Tuesday: next Monday
    assert next_run("week", [1], 9, 0, at("2026-10-27 10:00")) == at("2026-11-03 09:00")  # its own day, past: a week on
    assert next_run("week", [0, 3], 9, 0, at("2026-10-27 10:00")) == at("2026-10-29 09:00")  # Tuesday: then Thursday
    assert next_run("week", [0, 1, 2, 3, 4], 9, 0, at("2026-10-31 10:00")) == at("2026-11-02 09:00")  # Saturday: Monday
    assert next_run("week", [5, 6], 9, 0, at("2026-11-01 10:00")) == at("2026-11-07 09:00")  # Sunday, past: Saturday
    assert next_run("month", [5], 8, 0, at("2026-10-04 23:00")) == at("2026-10-05 08:00")
    assert next_run("month", [5], 8, 0, at("2026-12-05 08:00")) == at("2027-01-05 08:00")  # past December
    assert next_run("day", [], 8, 30, at("2026-10-27 08:10")) == at("2026-10-27 08:30")  # its minute, not the hour


@pytest.mark.communication
def test_a_schedule_said_in_plain_words():
    assert routines.when("week", [0, 1, 2, 3, 4], 8, 0) == "en semaine (du lundi au vendredi) à 8 h"
    assert routines.when("week", [5, 6], 8, 0) == "le week-end à 8 h"
    assert routines.when("week", [0, 3, 4], 9, 0) == "chaque lundi, jeudi et vendredi à 9 h"
    assert routines.when("week", [2], 9, 0) == "chaque mercredi à 9 h"
    assert routines.when("week", list(range(7)), 9, 0) == "chaque jour à 9 h"
    assert routines.when("month", [5], 8, 0) == "chaque mois, le 5 à 8 h"
    assert routines.when("week", [5, 6], 8, 5) == "le week-end à 8 h 05"
    assert "59 minutes" in routines.routine_problem("tva", "day", [], 8, 60, "")
    assert "quels jours" in routines.routine_problem("tva", "week", [], 8, 0, "")
    assert "un seul" in routines.routine_problem("tva", "month", [5, 6], 8, 0, "")
    assert "pas de jour" in routines.routine_problem("tva", "day", [99], 8, 0, "")


@pytest.fixture
def chat(db_cabinet, monkeypatch):
    """The chat (no sender to limit); the conversation's words; addresses @cabinet.ma allowed and live."""
    said = {"words": "chaque mois le 5 a 8h, rapport a chef@cabinet.ma"}
    monkeypatch.setattr(routines, "sender_reach", lambda session_id: None)
    monkeypatch.setattr(routines, "sender_of", lambda session_id: None)
    monkeypatch.setattr(routines, "recipient_allowed", lambda address: address.endswith("@cabinet.ma"))
    monkeypatch.setattr(routines, "live", lambda address: address != "parti@cabinet.ma")
    monkeypatch.setattr(routines, "client_messages", lambda session_id, message: [said["words"]])
    monkeypatch.setattr(routines, "confirmed", lambda session_id, yes, *parts: yes)
    return said


def make(yes, email="chef@cabinet.ma", hour=8, kind="tva", prompt=""):
    return routines.manage_routine(
        "create", kind, "month", [5], hour, email=email, prompt=prompt, client_confirmed=yes, session_id="s"
    )


@pytest.mark.security
def test_a_routine_is_read_back_then_made_on_the_clients_yes(chat):
    asked = make(False)
    assert asked["status"] == "question" and "Vous confirmez" in asked["message"]
    assert "préparer la déclaration de TVA de la dernière période, chaque mois, le 5 à 8 h" in asked["message"]
    assert "rapport à `chef@cabinet.ma`" in asked["message"] and "sur tous les dossiers du cabinet" in asked["message"]
    with cabinet_db() as db:
        assert db.execute("select count(*) from tarik_routines").fetchone() == (0,)
    done = make(True)
    assert done["status"] == "ok" and done["message"].startswith("Routine n° ")
    listed = routines.manage_routine("list")
    assert "1 routine(s)" in listed["message"] and "prochaine le" in listed["message"]


@pytest.mark.security
def test_a_report_address_not_written_not_allowed_or_not_live_is_asked(chat):
    assert make(True, email="autre@ailleurs.ma")["message"] == "À quelle adresse e-mail envoyer ses rapports ?"
    chat["words"] += " ou moi@gmail.com ou parti@cabinet.ma"
    assert "pas dans les adresses autorisées" in make(True, email="moi@gmail.com")["message"]
    assert "ni un administrateur actif du cabinet" in make(True, email="parti@cabinet.ma")["message"]


@pytest.mark.functionality
def test_a_declaration_routine_at_midnight_is_asked_another_hour(chat):
    assert "de 1 à 23" in make(True, hour=0)["message"]  # midnight in Morocco: the day before on the server


@pytest.mark.security
@pytest.mark.parametrize(("rank", "made"), [(2, False), (3, True)])
def test_only_admins_and_staff_make_or_resume_a_routine_and_each_manages_their_own(chat, monkeypatch, rank, made):
    make(True)  # made in the chat
    reach = {"id": 7, "name": "Chef", "company_ids": [5], "rank": rank}
    monkeypatch.setattr(routines, "sender_reach", lambda session_id: reach)
    monkeypatch.setattr(routines, "sender_of", lambda session_id: "chef@cabinet.ma")
    result = make(True)
    assert (result["status"] == "ok") is made  # an approver may not: a routine covers the whole cabinet
    if not made:
        assert result["status"] == "stop" and "administrateurs" in result["message"]
    shown = routines.manage_routine("list", session_id="s")
    assert ("1 routine(s)" in shown["message"]) is made  # the chat's routine is never theirs
    with cabinet_db() as db:
        (other,) = db.execute("select min(id) from tarik_routines").fetchone()
    assert "Quelle routine" in routines.manage_routine("pause", routine_id=other, session_id="s")["message"]


@pytest.mark.functionality
def test_a_routine_paused_resumed_and_deleted_after_a_yes(chat):
    make(True)
    with cabinet_db() as db:
        (routine_id,) = db.execute("select id from tarik_routines").fetchone()
    routines.manage_routine("pause", routine_id=routine_id)
    assert "(en pause)" in routines.manage_routine("list")["message"]
    assert "reprise : prochaine le" in routines.manage_routine("resume", routine_id=routine_id)["message"]
    assert "Je la supprime" in routines.manage_routine("delete", routine_id=routine_id)["message"]
    routines.manage_routine("delete", routine_id=routine_id, client_confirmed=True)
    assert routines.manage_routine("list")["routine"]


@pytest.mark.functionality
def test_an_instruction_routine_needs_its_words(chat):
    assert "quelle consigne" in make(True, kind="instruction")["message"]
    said = "Dis-moi les factures en brouillon."
    asked = make(False, kind="instruction", prompt=said)
    assert "suivre sa consigne « Dis-moi les factures en brouillon. »" in asked["message"]
    assert make(True, kind="instruction", prompt=said)["status"] == "ok"
    with cabinet_db() as db:
        assert db.execute("select kind, prompt from tarik_routines").fetchone() == ("instruction", said)


def due_routine(kind="progress", prompt=""):
    """Make a routine due an hour ago, its reports to chef@cabinet.ma; return its id."""
    with cabinet_db() as db:
        return db.execute(
            "insert into tarik_routines (kind, every, days, hour, notify, created_by, next_run_at, prompt) "
            "values (%s, 'day', '{}', 8, 'chef@cabinet.ma', 'chat', %s, %s) returning id",
            (kind, datetime.now(ZONE) - timedelta(hours=1), prompt),
        ).fetchone()[0]


@pytest.mark.functionality
def test_a_due_routine_is_claimed_once_and_moved_to_its_next_time(db_cabinet):
    routine_id = due_routine()
    due = clock.claimed()
    assert [row[0] for row in due] == [routine_id]
    assert clock.claimed() == due  # never reported (the server stopped): run again, the same period
    clock.record(routine_id, due[0][7], "fait")
    assert clock.claimed() == []  # reported: done
    with cabinet_db() as db:
        (upcoming,) = db.execute("select next_run_at from tarik_routines where id = %s", (routine_id,)).fetchone()
        db.execute("update tarik_routines set next_run_at = %s where id = %s", (due[0][7], routine_id))
    assert upcoming > datetime.now(ZONE) and upcoming.astimezone(ZONE).hour == 8
    assert clock.claimed() == []  # the same period due again: recorded, never run twice


@pytest.mark.functionality
def test_a_period_runs_once_here_whatever_changes_meanwhile_and_each_minute_has_its_own_conversation(monkeypatch):
    begun = []

    async def ran(cabinet, row):
        begun.append(row)

    monkeypatch.setattr(clock, "ran", ran)
    due = datetime(2026, 10, 27, 8, 0, tzinfo=ZONE)
    row = (3, "progress", "day", [], 8, 0, "chef@cabinet.ma", due, "", "")

    async def twice():
        clock.started("c", row)
        clock.started("c", (*row[:6], "autre@cabinet.ma", *row[7:]))  # its address changed while it runs
        await asyncio.gather(*clock.PENDING)

    asyncio.run(twice())
    assert begun == [row]
    assert routines.session_of(3, due) != routines.session_of(3, due.replace(minute=30))  # moved to 8:30 that day


@pytest.fixture
def running(db_cabinet, monkeypatch):
    """A routine's period run as its Routine workflow on the test cabinet: dossiers 5 and 6 opened by Tarik, its
    address live; the tools called and the workflows run nested."""
    state = {"calls": [], "nested": [], "files": {}}
    monkeypatch.setattr(flow, "opened_dossiers", lambda: [{"id": 5, "name": "ALPHA"}, {"id": 6, "name": "BETA"}])
    monkeypatch.setattr(flow, "address_problem", lambda notify: "")
    monkeypatch.setattr(flow, "keep_file", lambda session, name, data: state["files"].update({name: data}))

    async def call(cabinet, tool, **arguments):
        state["calls"].append((tool, arguments))
        if tool == "dossier_progress_figures":
            dossier_id = arguments["dossier_id"]
            figures = {"progress": 10 * dossier_id, "suspense": dossier_id - 5, "pieces": 4, "exercice": "2026"}
            return {"status": "ok", "message": f"Dossier n° {dossier_id}\nAvancement", "figures": figures}
        if tool == "send_email":
            return {"status": "ok", "message_id": f"<rapport-{len(state['calls'])}@tarik>"}
        return {"status": "ok"}

    async def nested(name, given, cabinet, session, node=None):
        state["nested"].append((name, given, session))
        if name == "discussion":
            return {"reply": f"fait : {given['message']}"}
        return {"message": f"TVA du dossier {given['call']['dossier_id']} préparée.", "outcomes": [{"status": "ok"}]}

    async def start(name, payload, slots=None, unattended=False):  # the trigger's run, as is (no run store here)
        state["unattended"] = unattended
        output = await graphs.COMPILED[name].ainvoke(payload["input"], {"run_id": "run"})
        done = asyncio.get_running_loop().create_future()
        done.set_result((output, None))
        return "run", done

    monkeypatch.setattr(flow, "call", call)
    monkeypatch.setattr(calls, "call", call)
    monkeypatch.setattr(clock.runs, "nested", nested)
    monkeypatch.setattr(clock.runs, "start", start)
    return state


@pytest.mark.functionality
def test_a_routine_works_on_every_dossier_tarik_opens_and_its_report_goes_to_its_address(running, db_cabinet):
    routine_id = due_routine("tva")
    (row,) = clock.claimed()
    report = asyncio.run(clock.ran(db_cabinet, row))
    assert f"Voici le rapport « Routine n° {routine_id} : préparer" in report
    assert "- **2** dossiers parcourus : **2** préparées" in report
    assert "**À voir**" not in report and "**En échec**" not in report  # only what happened is said
    assert [given["call"]["routine"] for _, given, _ in running["nested"]] == [True, True]  # made, never taken up
    (sent,) = [arguments for tool, arguments in running["calls"] if tool == "send_email"]
    assert sent["to"] == "chef@cabinet.ma" and sent["body"] == report
    assert sent["attach_session"].startswith(f"routine-{routine_id}-")  # the XML the runs made, and the report's
    with cabinet_db() as db:  # a reply is answered in the routine's own conversation, the report its first email
        thread = db.execute("select client_address, state, session_id from tarik_mail_threads").fetchone()
    assert thread == ("chef@cabinet.ma", "done", sent["attach_session"])
    sheet = openpyxl.load_workbook(io.BytesIO(next(d for n, d in running["files"].items() if n.endswith(".xlsx"))))
    rows = list(sheet.active.values)
    assert rows == [
        ("Dossier", "Résultat", "Détail"),
        ("ALPHA", "Préparée", "TVA du dossier 5 préparée."),
        ("BETA", "Préparée", "TVA du dossier 6 préparée."),
    ]
    text = next(d for n, d in running["files"].items() if n.endswith(".csv")).decode("utf-8-sig")
    assert text.splitlines()[:2] == ["Dossier;Résultat;Détail", "ALPHA;Préparée;TVA du dossier 5 préparée."]
    with cabinet_db() as db:
        assert db.execute("select report from tarik_routine_runs").fetchone() == (report,)


@pytest.mark.functionality
def test_an_instruction_is_handed_to_the_agent_once_in_the_routines_own_session(running, db_cabinet):
    routine_id = due_routine("instruction", "Liste les brouillons.")
    (row,) = clock.claimed()
    assert asyncio.run(clock.ran(db_cabinet, row)) == "fait : Liste les brouillons."
    ((name, _, session),) = running["nested"]
    assert name == "discussion" and session.startswith(f"routine-{routine_id}-")
    assert running["unattended"] is True  # its period runs unattended, and so does everything nested in it


@pytest.mark.security
def test_a_routine_whose_address_is_gone_is_paused_and_sends_nothing(running, db_cabinet, monkeypatch):
    routine_id = due_routine()
    (row,) = clock.claimed()
    monkeypatch.setattr(flow, "address_problem", lambda notify: "`chef@cabinet.ma` est parti")
    assert asyncio.run(clock.ran(db_cabinet, row)) == "Routine en pause : `chef@cabinet.ma` est parti."
    assert [tool for tool, _ in running["calls"] if tool not in ("save_files", "record_step")] == []
    with cabinet_db() as db:
        assert db.execute("select enabled from tarik_routines where id = %s", (routine_id,)).fetchone() == (False,)


@pytest.mark.communication
def test_a_failing_dossier_reads_as_tarik_says_never_as_the_error(running, db_cabinet, monkeypatch):
    due_routine("tva")
    (row,) = clock.claimed()

    async def nested(name, given, cabinet, session, node=None):
        raise ConnectionError("postgres://tarik:secret@db/tarik refused")

    monkeypatch.setattr(clock.runs, "nested", nested)
    report = asyncio.run(clock.ran(db_cabinet, row))
    assert "- **2** dossiers parcourus : **2** en échec" in report
    assert "**En échec**\n\n- `ALPHA`\n- `BETA`" in report
    texts = [data for name, data in running["files"].items() if not name.endswith(".xlsx")]
    assert all(b"secret" not in data and b"echec technique" in data for data in texts)


@pytest.mark.communication
def test_experio_unreachable_sends_nothing_and_says_it_as_tarik_does(running, db_cabinet, monkeypatch):
    due_routine("progress")
    (row,) = clock.claimed()

    def down(notify):
        raise ConnectionError("odoo://tarik:secret@experio refused")

    monkeypatch.setattr(flow, "address_problem", down)
    report = asyncio.run(clock.ran(db_cabinet, row))
    assert "secret" not in report and not [tool for tool, _ in running["calls"] if tool == "send_email"]


@pytest.mark.communication
def test_a_dossier_s_own_refusal_is_its_line_and_no_dossier_is_nothing_to_say(running, db_cabinet, monkeypatch):
    due_routine("progress")
    (row,) = clock.claimed()

    async def call(cabinet, tool, **arguments):
        if tool == "dossier_progress_figures":
            raise calls.ToolFailure("Dossier `ALPHA` : Experio n'a pas pu donner son état.")
        return {"status": "ok"}

    monkeypatch.setattr(flow, "call", call)
    assert "**2** en échec" in asyncio.run(clock.ran(db_cabinet, row))
    plain = next(data for name, data in running["files"].items() if name.endswith(".txt")).decode()
    assert plain.startswith("Dossier `ALPHA` : Experio n'a pas pu donner son état.")  # its own words
    monkeypatch.setattr(flow, "opened_dossiers", lambda: [])
    (row,) = [(*row[:7], row[7] - timedelta(days=1), *row[8:])]  # another period
    assert asyncio.run(clock.ran(db_cabinet, row)) == flow.NOTHING


@pytest.mark.functionality
def test_an_instruction_too_long_is_cut_off_and_said(running, db_cabinet, monkeypatch):
    due_routine("instruction", "Liste tout.")
    (row,) = clock.claimed()

    async def slow(name, given, cabinet, session, node=None):
        await asyncio.sleep(1)

    monkeypatch.setattr(clock.runs, "nested", slow)
    monkeypatch.setattr(flow, "INSTRUCTION_SECONDS", 0.01)
    assert asyncio.run(clock.ran(db_cabinet, row)).startswith("Consigne interrompue")


@pytest.mark.functionality
def test_the_dossiers_not_reached_before_the_next_period_are_left_to_it(running, db_cabinet):
    due_routine("progress")
    (row,) = clock.claimed()
    late = (*row[:7], datetime.now(ZONE) - timedelta(days=2), *row[8:])  # a period whose next one already came
    report = asyncio.run(clock.ran(db_cabinet, late))
    assert ": **2** non traités (la période suivante est arrivée avant)" in report
    assert not [tool for tool, _ in running["calls"] if tool == "dossier_progress_figures"]


@pytest.fixture
def experio_users(monkeypatch):
    """Experio's users as admins_where reads them: Chef an admin, Démis an admin demoted on one dossier by a
    « Profils par Dossier » row; Jane of Experio's staff; the domains asked."""
    asked = []

    def odoo_call(model, method, domain, **kw):
        asked.append(domain)
        if model == "exp.user.company.profile":
            return [{"user_id": [2, "Démis"]}]
        admins = [{"id": 1, "name": "Chef", "login": "chef@cabinet.ma"}]
        admins.append({"id": 2, "name": "Démis", "login": "demis@cabinet.ma"})
        logins = [d[2].lower() for d in domain if d[0] == "login"]
        return [user for user in admins if not logins or user["login"] in logins]

    monkeypatch.setattr(routines, "odoo_call", odoo_call)
    staff = {"jane@experio.ma": {"employee_type": "internal"}, "junior@cabinet.ma": {"employee_type": "employee"}}
    monkeypatch.setattr(routines, "collaborator", staff.get)
    return asked


@pytest.mark.security
def test_only_a_live_admin_on_every_dossier_or_staff_gets_reports_of_every_dossier(experio_users):
    assert routines.live("chef@cabinet.ma") and routines.live("jane@experio.ma")
    assert not routines.live("demis@cabinet.ma")  # lower on a dossier: Experio shows them less there
    assert not routines.live("junior@cabinet.ma") and not routines.live("parti@cabinet.ma")
    rule = experio_users[0]
    assert ("is_agent", "=", False) in rule  # Experio's technical accounts left out
    assert ("employee_state", "=", "active") in rule
    assert ("exp_profile_id.management_group", "=", "admin") in rule


@pytest.mark.functionality
def test_the_cabinets_admins_offered_are_the_same_rule_among_allowed_addresses(experio_users, monkeypatch):
    monkeypatch.setattr(routines, "allowed_recipients", lambda: ["@cabinet.ma"])
    assert routines.admins() == [{"name": "Chef", "email": "chef@cabinet.ma"}]


@pytest.mark.security
@pytest.mark.parametrize(
    ("configurable", "kept"),
    [
        ({"unattended": True}, False),  # a routine's period: nobody there to say yes
        ({}, True),  # a client in the chat
        ({"session_id": "routine-3-202610270800"}, True),  # a client's reply in a routine's conversation: attended
    ],
)
def test_a_routines_run_never_confirms_declares_nor_posts(monkeypatch, configurable, kept):
    from types import SimpleNamespace

    from engine import agent, runs

    handed = []

    async def nested(name, given, cabinet, session_id, node=None):
        handed.append(given["call"])
        return {"message": "fait"}

    monkeypatch.setattr(runs, "nested", nested)
    given = {"session_id": "chat-" + "0" * 32, "cabinet": "c"} | configurable
    runtime = SimpleNamespace(state={"messages": []}, config={"configurable": given})
    asked = {"client_confirmed": True, "declarer": True, "comptabiliser": "oui"}
    asyncio.run(agent.run_workflow("declarations", runtime, asked))
    (call,) = handed
    said = (call["client_confirmed"], call["declarer"], call["comptabiliser"] == "oui")
    assert said == (kept, kept, kept) and call.get("routine", False) is not kept


@pytest.mark.security
@pytest.mark.parametrize(("workflow", "runs_it"), [("declarations", True), ("export", True), ("dossier", False)])
def test_a_routines_run_starts_only_what_prepares(monkeypatch, workflow, runs_it):
    from types import SimpleNamespace

    from engine import agent, runs

    async def nested(name, given, cabinet, session_id, node=None):
        return {"message": f"{name} fait"}

    monkeypatch.setattr(runs, "nested", nested)
    session = {"session_id": "routine-3-202610270800", "cabinet": "c", "unattended": True}
    runtime = SimpleNamespace(state={"messages": []}, config={"configurable": session})
    answer = asyncio.run(agent.run_workflow(workflow, runtime, {}))
    assert (answer == f"{workflow} fait") is runs_it and (answer == agent.NOT_IN_ROUTINE) is not runs_it


@pytest.mark.security
def test_an_approver_resumes_no_routine_either(chat, monkeypatch):
    make(True)
    with cabinet_db() as db:
        db.execute("update tarik_routines set created_by = 'chef@cabinet.ma', enabled = false")
        (routine_id,) = db.execute("select id from tarik_routines").fetchone()
    monkeypatch.setattr(routines, "sender_reach", lambda session_id: {"id": 7, "company_ids": [], "rank": 2})
    monkeypatch.setattr(routines, "sender_of", lambda session_id: "chef@cabinet.ma")
    assert routines.manage_routine("resume", routine_id=routine_id, session_id="s")["status"] == "stop"
    assert routines.manage_routine("pause", routine_id=routine_id, session_id="s")["status"] == "ok"  # stopping is


@pytest.mark.communication
def test_a_progress_report_is_a_few_lines_and_every_dossier_s_figures_are_in_the_files(running, db_cabinet):
    due_routine("progress")
    (row,) = clock.claimed()
    report = asyncio.run(clock.ran(db_cabinet, row))
    excel, spreadsheet, text = running["files"]
    assert report.split("\n\n**En bref**\n\n")[1].splitlines() == [
        "- **2** dossiers parcourus : **2** lus",
        "- Avancement moyen des dossiers (celui de chacun selon Experio) : **55 %**",
        "- Opérations bancaires en suspens (comptabilisées, sans facture lettrée, exercice en cours) : **1** opération "
        "dans **1** dossier",
        "",
        "**Les moins avancés**",
        "",
        "1. `ALPHA` : 50 %",
        "2. `BETA` : 60 %",
        "",
        "**Le détail de chaque dossier est joint**",
        "",
        f"- `{excel}` : le tableau, un dossier par ligne, avec des filtres",
        f"- `{spreadsheet}` : le même tableau, pour un autre logiciel",
        f"- `{text}` : le détail complet, à lire tel quel",
    ]
    assert [name.rsplit(".", 1)[1] for name in running["files"]] == ["xlsx", "csv", "txt"]
    sheet = openpyxl.load_workbook(io.BytesIO(next(d for n, d in running["files"].items() if n.endswith(".xlsx"))))
    header, alpha, _ = sheet.active.values
    assert header[:3] == ("Dossier", "Résultat", "Avancement (%)") and alpha[:3] == ("ALPHA", "Lu", 50)
    shown, page = render(report)  # as the client reads it: headings bold, lists, never a raw mark
    assert "<strong>En bref</strong>" in page and "<ol" in page and "**" not in shown and "**" not in plain(report)


@pytest.mark.communication
def test_many_dossiers_to_see_are_named_five_then_counted():
    parts = [{"dossier": f"D{i}", "result": report_files.TO_SEE, "detail": "?"} for i in range(8)]
    body = report_files.summary("tva", parts, "routine n° 1", ["a.xlsx", "a.csv", "a.txt"])
    assert "- `D4`\n- et **3** autres dans les fichiers joints" in body and "`D5`" not in body


@pytest.mark.functionality
def test_a_period_still_running_is_never_claimed_again_behind_it(db_cabinet):
    routine_id = due_routine()
    (row,) = clock.claimed()
    with cabinet_db() as db:  # its next time already came, the first period not reported yet
        db.execute("update tarik_routines set next_run_at = now() - interval '1 minute' where id = %s", (routine_id,))
    assert [r[7] for r in clock.claimed()] == [row[7]]  # the unfinished one only, no second period


@pytest.mark.security
def test_a_dossier_s_name_never_runs_as_a_formula_nor_breaks_the_files_or_the_email():
    parts = [{"dossier": "=HYPERLINK(1)", "result": report_files.FAILED, "detail": "a\x0cb `x` [lien](https://x.ma)"}]
    header, rows = report_files.table("tva", parts)
    assert rows == [["'=HYPERLINK(1)", "Échec", "ab `x` [lien](https://x.ma)"]]
    assert openpyxl.load_workbook(io.BytesIO(report_files.workbook(header, rows))).active["A2"].data_type == "s"
    body = report_files.summary("tva", [{**parts[0], "dossier": "X` [Payer](https://evil.ma) `Y"}], "routine", "abc")
    assert "<a href" not in inline_html(body)  # a backtick in a name never ends its code span


@pytest.mark.communication
def test_a_declaration_already_there_is_said_so_and_a_dossier_done_is_never_behind(running, db_cabinet, monkeypatch):
    due_routine("tva")
    (row,) = clock.claimed()

    async def nested(name, given, cabinet, session, node=None):
        return {"message": "Déjà déclarée dans Experio.", "outcomes": [{"status": "ok", "already": True}]}

    monkeypatch.setattr(clock.runs, "nested", nested)
    assert "**2** déjà faites" in asyncio.run(clock.ran(db_cabinet, row))
    done = {"result": report_files.READ, "detail": "", "figures": {"progress": 100, "suspense": 0}}
    body = report_files.summary("progress", [{**done, "dossier": "A"}], "routine", "abc")
    assert "**100 %**" in body and "Les moins avancés" not in body


@pytest.mark.communication
def test_a_named_routine_s_email_bears_its_name(running, db_cabinet):
    routine_id = due_routine("instruction", "Demande les priorités du jour.")
    with cabinet_db() as db:
        db.execute("update tarik_routines set name = 'Rappel Manager' where id = %s", (routine_id,))
    (row,) = clock.claimed()
    assert clock.given(db_cabinet, row)["call"]["subject"] == "Rappel Manager (chaque jour à 8 h)"


@pytest.mark.functionality
def test_a_routine_is_changed_in_the_chat_only_in_what_the_client_asked_once_they_said_yes(chat):
    make(True)
    with cabinet_db() as db:
        (routine_id, before) = db.execute("select id, next_run_at from tarik_routines").fetchone()
    asked = routines.manage_routine("update", routine_id=routine_id, name="Rappel TVA", session_id="s")
    assert asked["status"] == "question" and "« Rappel TVA »" in asked["message"]
    with cabinet_db() as db:  # read back only: nothing changed yet
        assert db.execute("select name from tarik_routines").fetchone() == ("",)
    changes = {"name": "Rappel TVA", "hour": 9, "minute": 30}
    done = routines.manage_routine("update", routine_id=routine_id, **changes, client_confirmed=True, session_id="s")
    assert done["status"] == "ok"
    with cabinet_db() as db:
        row = db.execute("select name, kind, every, days, hour, minute, next_run_at from tarik_routines").fetchone()
    assert row[:6] == ("Rappel TVA", "tva", "month", [5], 9, 30)  # the rest kept as it was
    assert row[6] != before and row[6].astimezone(ZONE).minute == 30  # a new time: its next run moves


@pytest.mark.functionality
def test_a_change_with_nothing_new_or_wrong_is_asked_again(chat):
    make(True)
    with cabinet_db() as db:
        (routine_id,) = db.execute("select id from tarik_routines").fetchone()
    assert "Rien de nouveau" in routines.manage_routine("update", routine_id=routine_id, session_id="s")["message"]
    weekly = {"routine_id": routine_id, "every": "week", "client_confirmed": True, "session_id": "s"}
    wrong = routines.manage_routine("update", **weekly)
    assert wrong["status"] == "question" and "quels jours" in wrong["message"]  # a month's 5 is no weekday list


@pytest.mark.security
def test_by_email_a_routine_s_reports_never_move_to_another_address(chat, monkeypatch):
    make(True)
    with cabinet_db() as db:
        db.execute("update tarik_routines set created_by = 'chef@cabinet.ma'")
        (routine_id,) = db.execute("select id from tarik_routines").fetchone()
    monkeypatch.setattr(routines, "sender_reach", lambda session_id: {"id": 7, "company_ids": [], "rank": 4})
    monkeypatch.setattr(routines, "sender_of", lambda session_id: "chef@cabinet.ma")
    moved = routines.manage_routine("update", routine_id=routine_id, email="autre@cabinet.ma", session_id="s")
    assert moved["status"] == "stop" and "page Routines" in moved["message"]


@pytest.mark.security
def test_the_chat_and_the_page_refuse_a_routine_whose_address_left_alike(chat, monkeypatch):
    make(True)
    (row,) = routines.routines_of_cabinet()
    monkeypatch.setattr(routines, "live", lambda address: False)  # chef@cabinet.ma left the cabinet
    asked = routines.manage_routine("update", routine_id=row["id"], hour=10, client_confirmed=True, session_id="s")
    assert asked["status"] == "question" and "chef@cabinet.ma" in asked["message"]
    values = {key: row[key] for key in routines.FIELDS} | {"hour": 10}
    assert "chef@cabinet.ma" in routines.problem_of(values)  # the page's check, the same one
    with cabinet_db() as db:
        db.execute("update tarik_routines set enabled = false")
    resumed = routines.manage_routine("resume", routine_id=row["id"], session_id="s")
    assert resumed["status"] == "question" and "Pour reprendre" in resumed["message"]  # resumed only once fixed


@pytest.mark.security
def test_an_unattended_run_marks_every_run_nested_in_it_and_nothing_else():
    from engine import runs

    assert "unattended" not in runs.configurable_of("c", "s", None)
    marked = runs.UNATTENDED.set(True)  # inside a routine's period, as execute sets it for its nested runs
    try:
        assert runs.configurable_of("c", "s", "s")["unattended"] is True
    finally:
        runs.UNATTENDED.reset(marked)
    assert "unattended" not in runs.configurable_of("c", "routine-3-202610270800", None)  # never from its id


@pytest.mark.functionality
def test_a_routine_s_report_is_tarik_s_own_turn_in_its_conversation(monkeypatch):
    from langgraph.checkpoint.memory import InMemorySaver

    from engine import agent

    monkeypatch.setattr(agent, "CHECKPOINTER", InMemorySaver())
    asyncio.run(agent.began("c", "routine-2-202610290900", "Rappel Manager", "Par quoi\u200b commencer ?"))
    config = {"configurable": {"thread_id": "routine-2-202610290900", "checkpoint_ns": ""}}
    saved = asyncio.run(agent.CHECKPOINTER.aget(config))["channel_values"]["messages"]
    assert [(message.type, message.text) for message in saved] == [
        ("human", "[Routine : Rappel Manager]"),  # the line the conversation opens on, nobody's words
        # Tarik's: a reply reads as its answer; Experio's data in it stays data
        ("ai", '<untrusted_data source="Experio" label="rapport_envoye">\nPar quoi commencer ?\n</untrusted_data>'),
    ]
    assert saved[0].name == agent.ROUTINE_NAME  # the routine's line, never the client's


@pytest.mark.security
def test_a_run_nested_in_an_unattended_one_is_built_unattended(monkeypatch):
    from engine import runs

    built = []

    class Graph:
        async def ainvoke(self, given, config):
            return config["configurable"]

    monkeypatch.setattr(runs, "pool", None)
    monkeypatch.setattr(runs, "graph_of", lambda name, configurable: built.append(configurable) or Graph())

    async def inside_a_period():
        marked = runs.UNATTENDED.set(True)
        try:
            return await runs.nested("agent", {}, "c", "routine-2-202610290900", thread="routine-2-202610290900")
        finally:
            runs.UNATTENDED.reset(marked)

    assert asyncio.run(inside_a_period())["unattended"] is True  # what the agent's tools read (run_workflow)
    assert "unattended" not in asyncio.run(runs.nested("agent", {}, "c", "routine-2-202610290900"))  # a reply's


@pytest.mark.functionality
def test_a_period_run_again_after_its_report_went_never_sends_it_twice(running, db_cabinet):
    due_routine("tva")
    (row,) = clock.claimed()
    asyncio.run(clock.ran(db_cabinet, row))
    with cabinet_db() as db:
        db.execute("update tarik_routine_runs set report = null")  # the server stopped before its report was kept
    (again,) = clock.claimed()
    asyncio.run(clock.ran(db_cabinet, again))
    assert len([tool for tool, _ in running["calls"] if tool == "send_email"]) == 1


@pytest.mark.security
def test_an_instruction_cut_off_never_joins_its_conversation_to_a_reply(running, db_cabinet, monkeypatch):
    due_routine("instruction", "Liste tout.")
    (row,) = clock.claimed()

    async def slow(name, given, cabinet, session, node=None):
        await asyncio.sleep(1)

    monkeypatch.setattr(clock.runs, "nested", slow)
    monkeypatch.setattr(flow, "INSTRUCTION_SECONDS", 0.01)
    asyncio.run(clock.ran(db_cabinet, row))
    with cabinet_db() as db:
        assert db.execute("select count(*) from tarik_mail_threads").fetchone() == (0,)  # a tool call left open there


@pytest.mark.security
def test_a_routine_s_message_is_never_the_client_s_and_its_mark_never_reaches_the_model():
    from types import SimpleNamespace

    from langchain_anthropic.chat_models import _format_messages
    from langchain_core.messages import AIMessage, HumanMessage

    from engine import agent

    began = HumanMessage("[Routine : Rappel Manager]", name=agent.ROUTINE_NAME)
    history = [began, AIMessage("Par quoi commencer ?")]
    assert agent.client_message(SimpleNamespace(state={"messages": history})) == ""  # nobody wrote yet
    history.append(HumanMessage("Commence par ALPHA."))
    assert agent.client_message(SimpleNamespace(state={"messages": history})) == "Commence par ALPHA."
    _, sent = _format_messages(history, model=None)
    assert [message["role"] for message in sent] == ["user", "assistant", "user"]  # the Messages format, user first
    assert all("name" not in message for message in sent)  # the mark stays in Tarik's memory
